"""Read user-exported patent records locally; never contact a patent website.

The returned dictionaries preserve original field names and values.  Normalization
owns field aliases; this module owns file formats and source row coordinates.
XLSX formulas are returned as text and are never executed.
"""
from __future__ import annotations

import csv
import io
import json
from datetime import date, datetime
from pathlib import Path


def _json_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def _rows(headers, rows, *, first_row=2, sheet=None):
    names = [str(v).strip() if v is not None else "" for v in headers]
    if len(names) != len(set(names)) or any(not name for name in names):
        raise ValueError("导入表头必须非空且不能重复；请先明确每列名称。")
    if any(name.startswith("_source_") for name in names):
        raise ValueError("表头不能使用程序保留的 _source_ 前缀。")
    result = []
    next_line = first_row
    for values in rows:
        line = next_line
        next_line = getattr(rows, "line_num", line) + 1
        values = list(values)
        if not any(value is not None and str(value).strip() for value in values):
            continue
        if len(values) > len(names):
            if any(v is not None and str(v).strip() for v in values[len(names):]):
                raise ValueError(f"第 {line} 行的有效列数超过表头，无法可靠关联字段。")
            values = values[:len(names)]
        row = {name: _json_value(values[i]) if i < len(values) else None
               for i, name in enumerate(names)}
        row["_source_row"] = line
        if sheet is not None:
            row["_source_sheet"] = sheet
        result.append(row)
    return result


def _decode(data: bytes) -> str:
    # Both encodings are common in Chinese spreadsheet exports.  No lossy fallback.
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass
    raise ValueError("无法解码导出文件；请保存为 UTF-8 CSV 或 JSON。")


def read_export(path: Path) -> list[dict]:
    """Read CSV/TSV/JSON/XLSX; empty data means zero hits, invalid files raise.

    JSON accepts a list of objects or an object with a ``records`` list.  XLSX
    reads all nonempty sheets; source sheet and row remain attached to each row.
    No record is silently skipped on a malformed structure.
    """
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"找不到导出文件：{path}")
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        raw = _decode(path.read_bytes())
        if not raw.strip():
            return []
        delimiter = "\t" if suffix == ".tsv" else ","
        if suffix == ".csv":
            try:
                delimiter = csv.Sniffer().sniff(raw[:65536], delimiters=",\t;").delimiter
            except csv.Error:
                pass
        reader = csv.reader(io.StringIO(raw, newline=""), delimiter=delimiter, strict=True)
        headers = next(reader, [])
        return _rows(headers, reader)
    if suffix == ".json":
        payload = json.loads(_decode(path.read_bytes()))
        records = payload.get("records") if isinstance(payload, dict) else payload
        if not isinstance(records, list) or any(not isinstance(row, dict) for row in records):
            raise ValueError("JSON 必须为对象列表，或包含 records 对象列表。")
        result = []
        for index, record in enumerate(records, 1):
            if any(str(key).startswith("_source_") for key in record):
                raise ValueError("JSON 字段不能使用程序保留的 _source_ 前缀。")
            result.append({**record, "_source_row": index})
        return result
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("读取 XLSX 需要 openpyxl，请按 requirements.txt 安装依赖。") from exc
        workbook = load_workbook(path, read_only=True, data_only=False)
        result = []
        try:
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows, None)
                if headers is None or not any(v is not None for v in headers):
                    if any(any(v is not None for v in row) for row in rows):
                        raise ValueError(f"工作表 {sheet.title} 第一行缺少表头。")
                    continue
                # Spreadsheet formatting may extend beyond the real data table.
                headers = list(headers)
                while headers and headers[-1] is None:
                    headers.pop()
                result.extend(_rows(headers, rows, sheet=sheet.title))
        finally:
            workbook.close()
        return result
    raise ValueError(f"不支持 {suffix or '无扩展名'} 文件；请使用 CSV、TSV、JSON 或 XLSX。")
